"""
report_builder.py
=================
Part 5 report engine (no Streamlit code in this file, so it can be tested on its own).

exportable_report.py draws the interface; this module does the real work:

    generate_report(df, quality_results, analytics_results, recommendations, ...)
        -> a "content" dictionary holding every section of the report
    export_report(report_format, content, csv_table)
        -> (file_bytes, file_name, mime_type) for PDF, Excel or CSV

Design rules
------------
* The DataFrame is only READ, never modified.
* Every number in the report is computed from the user's data (or supplied by
  an earlier part of the platform). Nothing is invented. When something is not
  available, the report says so in its "notes".
* The report contains summaries and findings, not dataset rows. Actual category
  values are only included when the user ticks the privacy checkbox.
* Text going into Excel/CSV is protected against formula injection (cells that
  start with = + - @ are prefixed with an apostrophe).
"""

import io
import re
import warnings
from datetime import datetime
from xml.sax.saxutils import escape as _xml_escape

import numpy as np
import pandas as pd

from platform_helpers import (detect_datetime_columns, detect_numeric_columns,
                              make_unique_columns)

# ---------------------------------------------------------------------------
# PUBLIC CONSTANTS (imported by exportable_report.py)
# ---------------------------------------------------------------------------
DEFAULT_REPORT_TITLE = "Data Quality & Analytics Report"
REPORT_FORMATS = ["PDF", "Excel", "CSV"]
CSV_TABLE_OPTIONS = [
    "Dataset Summary", "Quality Findings", "Missing Values", "Outliers",
    "Invalid Values", "Data Types", "Consistency Findings",
    "Numerical Summary", "Categorical Summary", "Correlations",
    "Recommendations",
]

# ---------------------------------------------------------------------------
# INTERNAL CONSTANTS
# ---------------------------------------------------------------------------
PRIORITY_ORDER = ["High", "Medium", "Low", "Informational"]
_PRIORITY_ALIASES = {"high": "High", "medium": "Medium", "med": "Medium",
                     "low": "Low", "informational": "Informational",
                     "info": "Informational"}
# Recommendation field names accepted from the Recommendations part (case and
# punctuation are ignored), mapped to the report's column names.
_REC_FIELD_ALIASES = {
    "priority": "Priority", "category": "Category", "issue": "Issue",
    "finding": "Finding", "recommendation": "Recommendation",
    "suggested action": "Suggested Action", "action": "Suggested Action",
    "affected columns": "Affected Columns", "columns": "Affected Columns",
}
_REC_BASE_COLUMNS = ["Priority", "Category", "Issue", "Finding",
                     "Recommendation", "Suggested Action"]

IQR_MULTIPLIER = 1.5
STRONG_CORRELATION = 0.70
MAX_PDF_TABLE_ROWS = 40
MAX_CHART_COLUMNS = 3
PLACEHOLDER_TOKENS = {"", "n/a", "na", "null", "none", "nan", "unknown", "?",
                      "-", "--", "missing", "#n/a"}
# Words in a column name that suggest the values should never be negative.
NON_NEGATIVE_WORDS = {"age", "price", "cost", "amount", "quantity", "qty",
                      "count", "weight", "height", "length", "distance",
                      "salary", "income", "duration", "fare", "population"}


# ---------------------------------------------------------------------------
# SMALL HELPERS
# ---------------------------------------------------------------------------
def _safe(fn, default, notes, label):
    """Run fn(); on failure return default and record a note (never crash)."""
    try:
        return fn()
    except Exception as exc:  # one failing section must not stop the report
        notes.append(f"{label} could not be calculated ({type(exc).__name__}: {exc}).")
        return default


def _is_text_series(series):
    return series.dtype == object or isinstance(series.dtype, pd.StringDtype)


def _text_values(series):
    """Non-missing values of a column that are actual strings."""
    values = series.dropna()
    return values[values.map(lambda v: isinstance(v, str))]


def _finite(series):
    """Non-missing, non-infinite values of a numeric column as floats."""
    values = pd.to_numeric(series, errors="coerce").astype(float)
    return values.replace([np.inf, -np.inf], np.nan).dropna()


def _name_suggests_non_negative(name):
    """True when a column name contains a word such as age, price or quantity."""
    split = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", "_", str(name))
    words = [w for w in re.split(r"[^a-z0-9]+", split.lower()) if w]
    return any(w in NON_NEGATIVE_WORDS or w.rstrip("s") in NON_NEGATIVE_WORDS
               for w in words)


def _shorten(text, limit=28):
    text = str(text)
    return text if len(text) <= limit else text[: limit - 1] + "\u2026"


def _mpl_text(text):
    """Escape $ so matplotlib does not try to render column names as math."""
    return str(text).replace("$", r"\$")


def _round_frame(frame, digits=4):
    """Round float columns so tables stay readable."""
    out = frame.copy()
    for column in out.columns:
        if pd.api.types.is_float_dtype(out[column]):
            out[column] = out[column].round(digits)
    return out


def _strength_label(r):
    value = abs(r)
    if value >= 0.9:
        return "Very strong"
    if value >= STRONG_CORRELATION:
        return "Strong"
    if value >= 0.5:
        return "Moderate"
    return "Weak"


def _as_text_list(obj):
    """Turn a supplied 'findings' value (str, list, dicts) into a list of strings."""
    if obj is None:
        return []
    if isinstance(obj, str):
        return [obj] if obj.strip() else []
    if isinstance(obj, pd.DataFrame):
        return [" | ".join(str(v) for v in row) for row in obj.itertuples(index=False)]
    if isinstance(obj, dict):
        return [f"{k}: {v}" for k, v in obj.items()]
    try:
        items = list(obj)
    except TypeError:
        return [str(obj)]
    texts = []
    for item in items:
        if isinstance(item, dict):
            texts.append("; ".join(f"{k}: {v}" for k, v in item.items()))
        elif item is not None and str(item).strip():
            texts.append(str(item))
    return texts


def _extract_score(quality_results):
    """Overall quality score from the Data Quality Report, as text, or None."""
    for key in ("Quality Score", "quality_score", "Overall Quality Score",
                "overall_score", "score", "overall"):
        if key in quality_results and quality_results[key] not in (None, ""):
            value = quality_results[key]
            if isinstance(value, dict):
                value = value.get("overall", value.get("score"))
            if value is None:
                continue
            if isinstance(value, (int, float, np.integer, np.floating)) \
                    and not isinstance(value, bool):
                return f"{float(value):g}"
            return str(value)
    return None


# ---------------------------------------------------------------------------
# RECOMMENDATIONS (supplied by Part 4)
# ---------------------------------------------------------------------------
def _normalize_recommendations(recommendations):
    """Return a clean, priority-sorted DataFrame, or None when nothing usable."""
    if recommendations is None:
        return None
    if isinstance(recommendations, pd.DataFrame):
        records = recommendations.to_dict("records")
    elif isinstance(recommendations, dict):
        records = [recommendations]
    else:
        try:
            records = list(recommendations)
        except TypeError:
            return None
    cleaned = []
    for record in records:
        if not isinstance(record, dict):
            continue
        row = {}
        for key, value in record.items():
            canonical = _REC_FIELD_ALIASES.get(
                re.sub(r"[\s_\-]+", " ", str(key)).strip().lower())
            if canonical and canonical not in row:
                row[canonical] = "" if value is None or (
                    isinstance(value, float) and np.isnan(value)) else str(value)
        if row:
            cleaned.append(row)
    if not cleaned:
        return None
    columns = list(_REC_BASE_COLUMNS)
    if any(r.get("Affected Columns") for r in cleaned):
        columns.append("Affected Columns")
    table = pd.DataFrame(cleaned).reindex(columns=columns).fillna("")
    table["Priority"] = table["Priority"].map(
        lambda p: _PRIORITY_ALIASES.get(str(p).strip().lower(), str(p).strip().title()))
    rank = {name: i for i, name in enumerate(PRIORITY_ORDER)}
    table["_rank"] = table["Priority"].map(lambda p: rank.get(p, len(rank)))
    table = table.sort_values("_rank", kind="stable").drop(columns="_rank")
    return table.reset_index(drop=True)


# ---------------------------------------------------------------------------
# SECTION BUILDERS (each reads the DataFrame and returns tables / numbers)
# ---------------------------------------------------------------------------
def _build_dtype_table(work):
    rows = []
    for column in work.columns:
        series = work[column]
        try:
            unique = int(series.nunique(dropna=True))
        except TypeError:  # unhashable cells such as lists
            unique = "n/a"
        rows.append({"Column": column, "Data Type": str(series.dtype),
                     "Non-Null Count": int(series.notna().sum()),
                     "Unique Values": unique})
    return pd.DataFrame(rows)


def _build_outlier_table(work, numeric_cols):
    rows = []
    for column in numeric_cols:
        values = _finite(work[column])
        if len(values) < 4:
            continue
        q1, q3 = values.quantile(0.25), values.quantile(0.75)
        iqr = q3 - q1
        if iqr == 0:
            continue
        lower, upper = q1 - IQR_MULTIPLIER * iqr, q3 + IQR_MULTIPLIER * iqr
        count = int(((values < lower) | (values > upper)).sum())
        if count:
            rows.append({"Column": column, "Q1": q1, "Q3": q3,
                         "Lower Fence": lower, "Upper Fence": upper,
                         "Potential Outliers": count,
                         "Outlier %": count / len(values) * 100})
    return _round_frame(pd.DataFrame(
        rows, columns=["Column", "Q1", "Q3", "Lower Fence", "Upper Fence",
                       "Potential Outliers", "Outlier %"]), 2)


def _build_invalid_table(work, numeric_cols, quality_results):
    rows, seen = [], set()

    def add(column, issue, count):
        key = (str(column), issue)
        if count > 0 and key not in seen:
            seen.add(key)
            rows.append({"Column": str(column), "Issue": issue, "Count": int(count)})

    for column in numeric_cols:
        series = work[column]
        values = pd.to_numeric(series, errors="coerce").dropna().astype(float)
        add(column, "Infinite values", int(np.isinf(values).sum()))
        if _name_suggests_non_negative(column):
            add(column, "Negative values (the column name suggests values "
                        "should not be negative)", int((values < 0).sum()))
    for column in work.columns:
        if _is_text_series(work[column]):
            text = _text_values(work[column]).str.strip().str.lower()
            add(column, "Placeholder text values (such as N/A, unknown or ?)",
                int(text.isin(PLACEHOLDER_TOKENS).sum()))

    # Invalid values already reported by the Data Quality Report, if supplied.
    reported = quality_results.get("invalid_values")
    if isinstance(reported, pd.DataFrame):
        reported = reported.to_dict("records")
    for item in reported or []:
        if isinstance(item, dict) and {"column", "issue", "count"} <= set(item):
            try:
                add(item["column"], str(item["issue"]), int(item["count"]))
            except (TypeError, ValueError):
                continue
    return pd.DataFrame(rows, columns=["Column", "Issue", "Count"])


def _build_consistency_table(work, categorical_cols, include_values):
    rows = []
    for column in categorical_cols:
        series = work[column]
        if not _is_text_series(series):
            continue
        text = _text_values(series)
        if text.empty:
            continue
        # Mixed Python types in one column (e.g. numbers and words).
        type_names = sorted(series.dropna().map(lambda v: type(v).__name__).unique())
        if len(type_names) > 1:
            rows.append({"Column": column, "Issue": "Mixed value types",
                         "Details": ", ".join(type_names)})
        # Leading / trailing spaces split one category into several.
        padded = int((text != text.str.strip()).sum())
        if padded:
            rows.append({"Column": column, "Issue": "Leading or trailing spaces",
                         "Details": f"{padded:,} value(s)"})
        # Names that match once case, spacing and punctuation are ignored.
        stripped = text.str.strip()
        normalized = stripped.str.lower().str.replace(r"[\W_]+", "", regex=True)
        pairs = pd.DataFrame({"raw": stripped, "norm": normalized})
        pairs = pairs[pairs["norm"] != ""].drop_duplicates()
        groups = pairs.groupby("norm")["raw"].agg(list)
        groups = groups[groups.map(len) > 1]
        if len(groups):
            detail = f"{len(groups):,} group(s) of names differ only by case, spacing or punctuation"
            if include_values:
                example = " / ".join(f'"{v}"' for v in groups.iloc[0][:3])
                detail += f" (for example {example})"
            rows.append({"Column": column, "Issue": "Possibly inconsistent category names",
                         "Details": detail})
        # Numbers or dates that are stored as text.
        sample = stripped.head(2000).str.replace(r"[,$%\s]", "", regex=True)
        if pd.to_numeric(sample, errors="coerce").notna().mean() >= 0.9:
            rows.append({"Column": column, "Issue": "Numbers possibly stored as text",
                         "Details": "At least 90% of sampled values look numeric"})
    return pd.DataFrame(rows, columns=["Column", "Issue", "Details"])


def _build_numeric_stats(work, numeric_cols):
    rows = []
    for column in numeric_cols:
        values = _finite(work[column])
        if values.empty:
            continue
        rows.append({
            "Column": column, "Count": int(len(values)), "Mean": values.mean(),
            "Median": values.median(),
            "Std Dev": values.std() if len(values) > 1 else np.nan,
            "Min": values.min(), "Max": values.max(),
            "Skewness": values.skew() if len(values) > 2 else np.nan})
    return _round_frame(pd.DataFrame(
        rows, columns=["Column", "Count", "Mean", "Median", "Std Dev", "Min",
                       "Max", "Skewness"]))


def _build_categorical_summary(work, categorical_cols, include_values):
    rows = []
    for column in categorical_cols:
        series = work[column]
        try:
            counts = series.dropna().astype(str).value_counts()
        except Exception:  # unusual cell contents
            continue
        non_null = int(counts.sum())
        row = {"Column": column, "Unique Values": int(len(counts)),
               "Missing %": round(float(series.isna().mean() * 100), 2),
               "Top Value Share %": (round(float(counts.iloc[0] / non_null * 100), 1)
                                     if non_null else np.nan)}
        if include_values:
            row["Most Frequent Value"] = str(counts.index[0]) if non_null else ""
        rows.append(row)
    columns = ["Column", "Unique Values", "Missing %", "Top Value Share %"]
    if include_values:
        columns.append("Most Frequent Value")
    return pd.DataFrame(rows, columns=columns)


def _correlation_matrix(work, numeric_cols, analytics_results):
    """Pearson matrix: the one supplied by the analytics part, or computed here."""
    supplied = analytics_results.get("correlation_matrix")
    if isinstance(supplied, pd.DataFrame) and not supplied.empty \
            and supplied.shape[0] == supplied.shape[1]:
        return supplied.apply(pd.to_numeric, errors="coerce")
    usable = [c for c in numeric_cols
              if _finite(work[c]).nunique() > 1 and work[c].notna().sum() >= 3]
    if len(usable) < 2:
        return None
    numeric = work[usable].apply(pd.to_numeric, errors="coerce")
    numeric = numeric.replace([np.inf, -np.inf], np.nan)
    return numeric.corr()


def _correlation_pairs(matrix):
    columns = ["Variable 1", "Variable 2", "Correlation (r)", "Strength"]
    if matrix is None:
        return pd.DataFrame(columns=columns)
    values = matrix.to_numpy(dtype=float)
    rows = []
    for i in range(values.shape[0]):
        for j in range(i + 1, values.shape[1]):
            r = values[i, j]
            if not np.isnan(r):
                rows.append({"Variable 1": str(matrix.index[i]),
                             "Variable 2": str(matrix.columns[j]),
                             "Correlation (r)": round(float(r), 3),
                             "Strength": _strength_label(r)})
    table = pd.DataFrame(rows, columns=columns)
    if table.empty:
        return table
    order = table["Correlation (r)"].abs().sort_values(ascending=False).index
    return table.loc[order].head(15).reset_index(drop=True)


def _build_datetime_summary(work, datetime_cols):
    rows = []
    for column in datetime_cols:
        series = work[column]
        if not pd.api.types.is_datetime64_any_dtype(series):
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                try:
                    series = pd.to_datetime(series, errors="coerce", format="mixed")
                except (TypeError, ValueError):
                    series = pd.to_datetime(series, errors="coerce")
        valid = series.dropna()
        if valid.empty:
            continue
        rows.append({"Column": column, "Earliest": str(valid.min().date()),
                     "Latest": str(valid.max().date()),
                     "Span (days)": int((valid.max() - valid.min()).days),
                     "Missing": int(series.isna().sum())})
    return pd.DataFrame(rows, columns=["Column", "Earliest", "Latest",
                                       "Span (days)", "Missing"])


# ---------------------------------------------------------------------------
# CHARTS (drawn with matplotlib's Figure class: no global state, no GUI)
# ---------------------------------------------------------------------------
def _figure_png(fig):
    buffer = io.BytesIO()
    fig.savefig(buffer, format="png", dpi=130, bbox_inches="tight")
    return buffer.getvalue()


def _build_charts(work, numeric_cols, categorical_cols, missing_df, matrix,
                  analytics_results, include_values, notes):
    """Return a list of {'title': str, 'png': bytes}."""
    try:
        from matplotlib.figure import Figure
    except ImportError:
        notes.append("Matplotlib is not installed, so no charts were created.")
        return []

    charts = []

    def add_chart(title, build):
        try:
            fig = Figure(figsize=(7, 3.6))
            build(fig)
            charts.append({"title": title, "png": _figure_png(fig)})
        except Exception as exc:  # a chart failing must not break the report
            notes.append(f"The chart '{title}' could not be drawn ({type(exc).__name__}).")

    # 1. Missing values by column (percent of rows).
    if not missing_df.empty:
        top = missing_df.head(15)

        def draw_missing(fig):
            ax = fig.subplots()
            ax.barh([_mpl_text(_shorten(c)) for c in top["Column"]][::-1],
                    top["Missing %"].tolist()[::-1], color="#1E40AF")
            ax.set_xlabel("Missing values (% of rows)")
            ax.set_title("Missing Values by Column")
        add_chart("Missing Values by Column", draw_missing)

    # 2. Histograms of the first few numeric columns that vary.
    drawn = 0
    for column in numeric_cols:
        values = _finite(work[column])
        if len(values) < 5 or values.nunique() < 2:
            continue
        if drawn >= MAX_CHART_COLUMNS:
            break
        drawn += 1

        def draw_hist(fig, column=column, values=values):
            ax = fig.subplots()
            bins = int(min(20, max(5, np.sqrt(len(values)))))
            ax.hist(values, bins=bins, color="#3B82F6", edgecolor="white")
            ax.set_xlabel(_mpl_text(column))
            ax.set_ylabel("Number of rows")
            ax.set_title(_mpl_text(f"Distribution of {_shorten(column, 40)}"))
        add_chart(f"Distribution of {column}", draw_hist)

    # 3. Correlation heatmap.
    if matrix is not None and matrix.shape[0] >= 2:
        small = matrix.iloc[:12, :12]

        def draw_corr(fig):
            ax = fig.subplots()
            image = ax.imshow(small.to_numpy(dtype=float), vmin=-1, vmax=1, cmap="RdBu_r")
            labels = [_mpl_text(_shorten(c, 14)) for c in small.columns]
            ax.set_xticks(range(len(labels)))
            ax.set_yticks(range(len(labels)))
            ax.set_xticklabels(labels, rotation=45, ha="right", fontsize=7)
            ax.set_yticklabels(labels, fontsize=7)
            if len(labels) <= 8:
                for i in range(len(labels)):
                    for j in range(len(labels)):
                        value = small.iloc[i, j]
                        if not np.isnan(value):
                            ax.text(j, i, f"{value:.2f}", ha="center", va="center", fontsize=7)
            fig.colorbar(image, ax=ax)
            ax.set_title("Correlation Heatmap")
        add_chart("Correlation Heatmap", draw_corr)

    # 4. Top categories (only when the user allowed actual data values).
    if include_values:
        shown = 0
        for column in categorical_cols:
            try:
                counts = work[column].dropna().astype(str).value_counts()
            except Exception:
                continue
            if not 2 <= len(counts) <= 30 or shown >= 2:
                continue
            shown += 1
            top = counts.head(8)

            def draw_cat(fig, column=column, top=top):
                ax = fig.subplots()
                ax.barh([_mpl_text(_shorten(v)) for v in top.index][::-1],
                        top.tolist()[::-1], color="#D97706")
                ax.set_xlabel("Number of rows")
                ax.set_title(_mpl_text(f"Most Frequent Values: {_shorten(column, 40)}"))
            add_chart(f"Most Frequent Values: {column}", draw_cat)

    # 5. Figures supplied by the Analytics part (matplotlib figures or PNG bytes).
    figures = analytics_results.get("figures")
    if isinstance(figures, dict):
        figures = list(figures.items())
    for item in figures or []:
        try:
            title, figure = item
            if isinstance(figure, (bytes, bytearray)):
                png = bytes(figure)
            elif hasattr(figure, "savefig"):
                png = _figure_png(figure)
            else:
                notes.append(f"The supplied chart '{title}' was skipped (unsupported type).")
                continue
            charts.append({"title": str(title), "png": png})
        except Exception as exc:
            notes.append(f"A supplied chart could not be used ({type(exc).__name__}).")
    return charts


# ---------------------------------------------------------------------------
# REPORT CONTENT
# ---------------------------------------------------------------------------
def generate_report(df, quality_results=None, analytics_results=None,
                    recommendations=None, report_title=DEFAULT_REPORT_TITLE,
                    include_charts=True, include_category_values=False):
    """
    Build every section of the report from the DataFrame (read-only) and the
    optional results of earlier parts. Returns the "content" dictionary used by
    the on-screen preview and by all three exporters.
    """
    if not isinstance(df, pd.DataFrame):
        raise ValueError("A pandas DataFrame is required to build a report.")
    if df.shape[0] == 0 or df.shape[1] == 0:
        raise ValueError("The dataset has no rows or no columns, so there is nothing to report on.")

    quality_results = quality_results if isinstance(quality_results, dict) else {}
    analytics_results = analytics_results if isinstance(analytics_results, dict) else {}
    notes = []
    title = re.sub(r"[\x00-\x1f]+", " ", str(report_title or "")).strip() or DEFAULT_REPORT_TITLE
    now = datetime.now()

    # Work on a shallow copy so renaming columns never touches the caller's data.
    work = df.copy(deep=False)
    unique_names, renamed = make_unique_columns(work.columns)
    work.columns = [str(c) for c in unique_names]
    if renamed:
        notes.append("Duplicate column names were renamed for analysis: "
                     + ", ".join(f"{old} -> {new}" for old, new in renamed[:8]) + ".")

    rows, cols = work.shape
    numeric_cols = detect_numeric_columns(work)
    datetime_cols = _safe(lambda: detect_datetime_columns(work), [], notes, "Date/time detection")
    categorical_cols = [c for c in work.columns
                        if c not in numeric_cols and c not in datetime_cols
                        and not pd.api.types.is_datetime64_any_dtype(work[c])]

    # ---- dataset summary --------------------------------------------------
    missing_counts = work.isna().sum()
    missing_total = int(missing_counts.sum())
    total_cells = rows * cols
    missing_df = pd.DataFrame({
        "Column": missing_counts.index,
        "Missing Count": missing_counts.values,
        "Missing %": (missing_counts.values / rows * 100).round(2)})
    missing_df = (missing_df[missing_df["Missing Count"] > 0]
                  .sort_values("Missing Count", ascending=False).reset_index(drop=True))
    try:
        duplicate_rows = int(work.duplicated().sum())
        duplicate_pct = round(duplicate_rows / rows * 100, 2)
    except TypeError:
        duplicate_rows = duplicate_pct = None
        notes.append("Duplicate rows could not be counted because some columns "
                     "contain unhashable values (such as lists).")
    memory_mb = _safe(lambda: work.memory_usage(deep=False).sum() / 1024 ** 2, None, notes, "Memory usage")
    summary_items = [
        ("Rows", f"{rows:,}"), ("Columns", f"{cols:,}"),
        ("Numerical columns", f"{len(numeric_cols):,}"),
        ("Categorical columns", f"{len(categorical_cols):,}"),
        ("Date/time columns", f"{len(datetime_cols):,}"),
        ("Missing values", f"{missing_total:,} ({missing_total / total_cells * 100:.2f}% of all cells)"),
        ("Duplicate rows", "N/A" if duplicate_rows is None else f"{duplicate_rows:,} ({duplicate_pct}% of rows)"),
    ]
    if memory_mb is not None:
        summary_items.append(("Approximate memory use", f"{memory_mb:,.2f} MB"))
    summary = {"rows": rows, "columns": cols, "missing_total": missing_total,
               "duplicate_rows": duplicate_rows,
               "table": pd.DataFrame(summary_items, columns=["Item", "Value"])}

    # ---- data quality -----------------------------------------------------
    score = _extract_score(quality_results)
    dtype_table = _safe(lambda: _build_dtype_table(work), pd.DataFrame(), notes, "Data types")
    outlier_table = _safe(lambda: _build_outlier_table(work, numeric_cols), pd.DataFrame(), notes, "Outliers")
    invalid_table = _safe(lambda: _build_invalid_table(work, numeric_cols, quality_results),
                          pd.DataFrame(), notes, "Invalid values")
    consistency_table = _safe(
        lambda: _build_consistency_table(work, categorical_cols, include_category_values),
        pd.DataFrame(), notes, "Consistency findings")

    findings = []
    if missing_total:
        worst = missing_df.iloc[0]
        findings.append(("Missing Values",
                         f"{missing_total:,} missing values across {len(missing_df)} column(s); "
                         f"the most affected column is '{worst['Column']}' ({worst['Missing %']}% missing)."))
    if duplicate_rows:
        findings.append(("Duplicate Rows", f"{duplicate_rows:,} duplicate row(s) ({duplicate_pct}% of rows)."))
    if len(outlier_table):
        findings.append(("Outliers", f"Potential outliers (IQR method) in {len(outlier_table)} numerical column(s)."))
    if len(invalid_table):
        findings.append(("Invalid Values", f"Possibly invalid values in {invalid_table['Column'].nunique()} column(s)."))
    if len(consistency_table):
        findings.append(("Consistency", f"{len(consistency_table)} consistency finding(s) in text columns."))
    findings_rows = [{"Check": a, "Finding": b, "Source": "Computed from the dataset"} for a, b in findings]
    findings_rows += [{"Check": "Data Quality Report", "Finding": text, "Source": "Supplied by the Data Quality Report"}
                      for text in _as_text_list(quality_results.get("findings"))]
    quality = {
        "score": score,
        "findings_table": pd.DataFrame(findings_rows, columns=["Check", "Finding", "Source"]),
        "missing_table": missing_df, "duplicate_rows": duplicate_rows, "duplicate_pct": duplicate_pct,
        "outlier_table": outlier_table, "invalid_table": invalid_table,
        "dtype_table": dtype_table, "consistency_table": consistency_table,
    }
    if score is None:
        notes.append("No overall quality score was supplied, so none is shown.")

    # ---- analytics --------------------------------------------------------
    matrix = _safe(lambda: _correlation_matrix(work, numeric_cols, analytics_results), None, notes, "Correlations")
    pairs = _safe(lambda: _correlation_pairs(matrix), pd.DataFrame(), notes, "Correlation pairs")
    strong_pairs = (pairs[pairs["Correlation (r)"].abs() >= STRONG_CORRELATION].reset_index(drop=True)
                    if len(pairs) else pairs)
    analytics = {
        "numeric_stats": _safe(lambda: _build_numeric_stats(work, numeric_cols), pd.DataFrame(), notes, "Numerical summary"),
        "categorical_summary": _safe(lambda: _build_categorical_summary(work, categorical_cols, include_category_values),
                                     pd.DataFrame(), notes, "Categorical summary"),
        "correlation_pairs": pairs, "strong_pairs": strong_pairs,
        "datetime_summary": _safe(lambda: _build_datetime_summary(work, datetime_cols), pd.DataFrame(), notes, "Date/time summary"),
        "supplied_findings": _as_text_list(analytics_results.get("findings")),
    }

    # ---- charts -----------------------------------------------------------
    charts = []
    if include_charts:
        charts = _safe(lambda: _build_charts(work, numeric_cols, categorical_cols, missing_df, matrix,
                                             analytics_results, include_category_values, notes),
                       [], notes, "Charts")

    # ---- recommendations --------------------------------------------------
    rec_table = _normalize_recommendations(recommendations)
    rec_available = rec_table is not None and len(rec_table) > 0
    if not rec_available:
        rec_table = pd.DataFrame(columns=_REC_BASE_COLUMNS)
        notes.append("No recommendation results were supplied, so the Recommendations section is not available.")
    recs = {"available": rec_available, "table": rec_table}

    # ---- executive summary and actions (built only from the findings above) --
    summary_lines = [f"The dataset contains {rows:,} rows and {cols:,} columns "
                     f"({len(numeric_cols)} numerical, {len(categorical_cols)} categorical, "
                     f"{len(datetime_cols)} date/time)."]
    summary_lines.append(
        f"{missing_total:,} missing values were found ({missing_total / total_cells * 100:.1f}% of all cells)."
        if missing_total else "No missing values were found.")
    if duplicate_rows is None:
        summary_lines.append("Duplicate rows could not be counted for this dataset.")
    else:
        summary_lines.append(f"{duplicate_rows:,} duplicate row(s) were found." if duplicate_rows
                             else "No duplicate rows were found.")
    if len(outlier_table):
        summary_lines.append(f"Potential outliers were flagged in {len(outlier_table)} numerical column(s); "
                             "they may be valid values.")
    if len(invalid_table):
        summary_lines.append(f"Possibly invalid values were flagged in {invalid_table['Column'].nunique()} column(s).")
    if len(consistency_table):
        summary_lines.append(f"{len(consistency_table)} consistency finding(s) were detected in text columns.")
    if len(strong_pairs):
        summary_lines.append(f"{len(strong_pairs)} strongly correlated variable pair(s) were found "
                             "(correlation is not causation).")
    if score:
        summary_lines.append(f"The Data Quality Report supplied an overall quality score of {score}.")
    if rec_available:
        counts = rec_table["Priority"].value_counts()
        summary_lines.append(
            f"{len(rec_table)} recommendation(s) were supplied: "
            + ", ".join(f"{int(counts.get(p, 0))} {p}" for p in PRIORITY_ORDER if counts.get(p, 0)) + ".")

    actions = []
    if rec_available:
        seen = set()
        for _, rec in rec_table.iterrows():
            action = str(rec["Suggested Action"]).strip()
            if action and action not in seen and rec["Priority"] != "Informational":
                seen.add(action)
                actions.append(f"{rec['Category']}: {action}")
            if len(actions) >= 6:
                break
    if not actions:
        if missing_total:
            actions.append(f"Review the {len(missing_df)} column(s) with missing values and decide whether to fill, collect or exclude them.")
        if duplicate_rows:
            actions.append("Review duplicate rows with the data owner before removing any of them.")
        if len(outlier_table):
            actions.append("Inspect potential outliers before removing them; extreme values can be valid.")
        if len(invalid_table):
            actions.append("Review the flagged invalid or placeholder values and correct them at the source.")
        if len(consistency_table):
            actions.append("Standardize inconsistent text values (spacing, capitalization, data types).")
    actions.append("Apply any changes to a copy of the data, then re-run the platform to confirm the issues are resolved."
                   if (actions or missing_total) else
                   "No immediate data-quality action is needed based on the checks performed.")

    return {
        "title": title, "generated_at": now,
        "generated_at_str": now.strftime("%B %d, %Y at %I:%M %p"),
        "executive_summary": summary_lines, "summary": summary, "quality": quality,
        "analytics": analytics, "charts": charts, "recommendations": recs,
        "actions": actions, "notes": notes,
    }


# ---------------------------------------------------------------------------
# EXPORT HELPERS
# ---------------------------------------------------------------------------
def _sanitize_text(value):
    """Prefix an apostrophe when text could be read as a spreadsheet formula."""
    if isinstance(value, str) and value and value[0] in ("=", "+", "-", "@", "\t", "\r"):
        try:
            float(value.replace(",", ""))
            return value  # an ordinary number such as "-5" is not a formula
        except ValueError:
            return "'" + value
    return value


def _safe_frame(frame):
    """Copy of a table with formula-safe headers and text cells."""
    out = frame.copy()
    out.columns = [_sanitize_text(str(c)) for c in out.columns]
    for column in out.columns:
        if out[column].dtype == object:
            out[column] = out[column].map(_sanitize_text)
    return out


def _tables_for_csv(content):
    quality, analytics = content["quality"], content["analytics"]
    return {
        "Dataset Summary": content["summary"]["table"],
        "Quality Findings": quality["findings_table"],
        "Missing Values": quality["missing_table"],
        "Outliers": quality["outlier_table"],
        "Invalid Values": quality["invalid_table"],
        "Data Types": quality["dtype_table"],
        "Consistency Findings": quality["consistency_table"],
        "Numerical Summary": analytics["numeric_stats"],
        "Categorical Summary": analytics["categorical_summary"],
        "Correlations": analytics["correlation_pairs"],
        "Recommendations": content["recommendations"]["table"],
    }


def _file_stem(content):
    slug = re.sub(r"[^A-Za-z0-9]+", "_", content["title"]).strip("_").lower()[:40] or "report"
    return f"{slug}_{content['generated_at']:%Y%m%d_%H%M}"


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------
def generate_csv_report(report_content, table_name=None):
    """One results table as UTF-8 CSV bytes (CSV files hold a single table)."""
    table_name = table_name or CSV_TABLE_OPTIONS[0]
    tables = _tables_for_csv(report_content)
    if table_name not in tables:
        raise ValueError(f"Unknown CSV table: {table_name!r}.")
    table = tables[table_name]
    if table is None or len(table) == 0:
        raise ValueError(f"The '{table_name}' table has no rows for this dataset, so there is nothing to export.")
    return _safe_frame(table).to_csv(index=False).encode("utf-8")


# ---------------------------------------------------------------------------
# EXCEL
# ---------------------------------------------------------------------------
def _excel_value(value):
    if value is None:
        return None
    if isinstance(value, (np.generic,)):
        value = value.item()
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, (pd.Timestamp, datetime)):
        value = str(value)
    if isinstance(value, (list, dict, tuple, set)):
        value = str(value)
    return _sanitize_text(value) if isinstance(value, str) else value


def generate_excel_report(report_content):
    """The full report as an .xlsx workbook (bytes): Summary sheet first, one sheet per table."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    content = report_content
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1E40AF")
    header_font = Font(bold=True, color="FFFFFF")

    def fit_columns(ws):
        for index, column_cells in enumerate(ws.columns, start=1):
            longest = max((len(str(c.value)) for c in column_cells if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(index)].width = min(max(10, longest + 2), 70)

    # ---- Summary sheet -----------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = _sanitize_text(content["title"])
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"Generated on {content['generated_at_str']}"
    row = 4
    ws.cell(row=row, column=1, value="Executive Summary").font = Font(bold=True, size=12)
    for line in content["executive_summary"]:
        row += 1
        ws.cell(row=row, column=1, value=_excel_value(line))
    row += 2
    ws.cell(row=row, column=1, value="Dataset Summary").font = Font(bold=True, size=12)
    row += 1
    for column_index, name in enumerate(content["summary"]["table"].columns, start=1):
        cell = ws.cell(row=row, column=column_index, value=name)
        cell.fill, cell.font = header_fill, header_font
    for _, item in content["summary"]["table"].iterrows():
        row += 1
        for column_index, value in enumerate(item, start=1):
            ws.cell(row=row, column=column_index, value=_excel_value(value))
    if content["quality"]["score"]:
        row += 2
        ws.cell(row=row, column=1, value="Overall Quality Score").font = Font(bold=True)
        ws.cell(row=row, column=2, value=_excel_value(content["quality"]["score"]))
    row += 2
    ws.cell(row=row, column=1, value="Recommended Actions").font = Font(bold=True, size=12)
    for number, action in enumerate(content["actions"], start=1):
        row += 1
        ws.cell(row=row, column=1, value=_excel_value(f"{number}. {action}"))
    if content["notes"]:
        row += 2
        ws.cell(row=row, column=1, value="Report Notes").font = Font(bold=True, size=12)
        for note in content["notes"]:
            row += 1
            ws.cell(row=row, column=1, value=_excel_value(note))
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 45

    # ---- one sheet per table (header in row 1, data from row 2) ------------
    sheet_tables = [
        ("Quality Findings", content["quality"]["findings_table"], "No data-quality findings."),
        ("Missing Values", content["quality"]["missing_table"], "No missing values were found."),
        ("Outliers", content["quality"]["outlier_table"], "No potential outliers were found."),
        ("Invalid Values", content["quality"]["invalid_table"], "No invalid values were found."),
        ("Data Types", content["quality"]["dtype_table"], "No column information."),
        ("Consistency", content["quality"]["consistency_table"], "No inconsistencies were detected."),
        ("Numerical Summary", content["analytics"]["numeric_stats"], "No numerical columns."),
        ("Categorical Summary", content["analytics"]["categorical_summary"], "No categorical columns."),
        ("Correlations", content["analytics"]["correlation_pairs"], "Not enough numerical columns."),
        ("Date-Time Summary", content["analytics"]["datetime_summary"], "No date/time columns."),
        ("Recommendations", content["recommendations"]["table"],
         "Recommendation results are not available."),
    ]
    for name, table, empty_message in sheet_tables:
        sheet = wb.create_sheet(name[:31])
        for column_index, header in enumerate(table.columns, start=1):
            cell = sheet.cell(row=1, column=column_index, value=_sanitize_text(str(header)))
            cell.fill, cell.font = header_fill, header_font
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if len(table) == 0:
            sheet.cell(row=2, column=1, value=empty_message)
        for row_index, record in enumerate(table.itertuples(index=False), start=2):
            for column_index, value in enumerate(record, start=1):
                sheet.cell(row=row_index, column=column_index, value=_excel_value(value))
        sheet.freeze_panes = "A2"
        fit_columns(sheet)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# PDF (ReportLab)
# ---------------------------------------------------------------------------
def generate_pdf_report(report_content):
    """The full report as a PDF (bytes), with tables, charts and 'Page X of Y' footers."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    from reportlab.platypus import (Image, KeepTogether, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    content = report_content
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=9.5, leading=13)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10, textColor=colors.HexColor("#475569"))
    cell_style = ParagraphStyle("cell", parent=body, fontSize=7.5, leading=9)
    head_style = ParagraphStyle("head", parent=cell_style, fontName="Helvetica-Bold", textColor=colors.white)
    h1 = ParagraphStyle("h1", parent=styles["Title"], fontSize=20, leading=24, textColor=colors.HexColor("#1E3A8A"))
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], textColor=colors.HexColor("#1E40AF"), spaceBefore=12)
    h3 = ParagraphStyle("h3", parent=styles["Heading3"], textColor=colors.HexColor("#1E3A8A"), spaceBefore=8)
    page_width = letter[0] - 1.5 * inch

    def esc(value):
        # Escape &, < and > (and drop control characters) so any text is safe in a Paragraph.
        text = re.sub(r"[\x00-\x08\x0b-\x1f]", " ", str(value))
        return _xml_escape(text).replace("\n", "<br/>")

    def fmt(value):
        if value is None or (isinstance(value, float) and np.isnan(value)):
            return ""
        if isinstance(value, (float, np.floating)):
            return f"{value:,.3f}".rstrip("0").rstrip(".")
        if isinstance(value, (int, np.integer)) and not isinstance(value, bool):
            return f"{int(value):,}"
        return str(value)

    def table_flowable(frame):
        data = [[Paragraph(esc(c), head_style) for c in frame.columns]]
        for record in frame.head(MAX_PDF_TABLE_ROWS).itertuples(index=False):
            data.append([Paragraph(esc(fmt(v)), cell_style) for v in record])
        widths = [page_width / len(frame.columns)] * len(frame.columns)
        table = Table(data, colWidths=widths, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E40AF")),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return table

    story = []

    def add_table(frame, empty_message):
        if frame is None or len(frame) == 0:
            story.append(Paragraph(esc(empty_message), small))
            return
        story.append(table_flowable(frame))
        if len(frame) > MAX_PDF_TABLE_ROWS:
            story.append(Paragraph(esc(f"Showing the first {MAX_PDF_TABLE_ROWS} of {len(frame):,} rows. "
                                       "Use the Excel or CSV export for the full table."), small))
        story.append(Spacer(1, 6))

    # ---- title and overview -----------------------------------------------
    story.append(Paragraph(esc(content["title"]), h1))
    story.append(Paragraph(esc(f"Generated on {content['generated_at_str']}"), small))
    story.append(Paragraph("Reports are generated locally. They contain summaries and findings, not dataset rows.", small))
    story.append(Paragraph("Executive Summary", h2))
    for line in content["executive_summary"]:
        story.append(Paragraph(esc(line), body, bulletText="\u2022"))

    story.append(Paragraph("Dataset Summary", h2))
    add_table(content["summary"]["table"], "No dataset summary is available.")

    # ---- data quality -----------------------------------------------------
    quality = content["quality"]
    story.append(Paragraph("Data Quality Summary", h2))
    if quality["score"]:
        story.append(Paragraph(esc(f"Overall Quality Score: {quality['score']}"), body))
    else:
        story.append(Paragraph("No overall quality score was supplied, so none is shown.", small))
    add_table(quality["findings_table"], "No data-quality findings are available.")
    story.append(Paragraph("Missing Values by Column", h3))
    add_table(quality["missing_table"], "No missing values were found.")
    story.append(Paragraph("Duplicate Rows", h3))
    if quality["duplicate_rows"] is None:
        story.append(Paragraph("Duplicate rows could not be counted for this dataset.", small))
    else:
        story.append(Paragraph(esc(f"{quality['duplicate_rows']:,} duplicate row(s) ({quality['duplicate_pct']}% of rows). "
                                   "Duplicates are reported only; nothing is deleted."), body))
    story.append(Paragraph("Potential Outliers (IQR method)", h3))
    story.append(Paragraph("These are statistical flags and are not automatically incorrect values.", small))
    add_table(quality["outlier_table"], "No potential outliers were found.")
    story.append(Paragraph("Invalid Values", h3))
    add_table(quality["invalid_table"], "No invalid values were found by the checks performed.")
    story.append(Paragraph("Data Types", h3))
    add_table(quality["dtype_table"], "No column information is available.")
    story.append(Paragraph("Consistency Findings", h3))
    add_table(quality["consistency_table"], "No categorical inconsistencies were detected.")

    # ---- analytics --------------------------------------------------------
    analytics = content["analytics"]
    story.append(Paragraph("Analytics Summary", h2))
    story.append(Paragraph("Numerical Summary", h3))
    add_table(analytics["numeric_stats"], "The dataset has no numerical columns to summarise.")
    story.append(Paragraph("Categorical Summary", h3))
    add_table(analytics["categorical_summary"], "The dataset has no categorical columns to summarise.")
    story.append(Paragraph("Correlation Analysis", h3))
    add_table(analytics["correlation_pairs"], "At least two numerical columns are needed for correlations.")
    if len(analytics["datetime_summary"]):
        story.append(Paragraph("Date/Time Columns", h3))
        add_table(analytics["datetime_summary"], "")
    if analytics["supplied_findings"]:
        story.append(Paragraph("Findings from the Analytics Section", h3))
        for text in analytics["supplied_findings"]:
            story.append(Paragraph(esc(text), body, bulletText="\u2022"))
    story.append(Paragraph("Charts", h3))
    if content["charts"]:
        for chart in content["charts"]:
            try:
                width_px, height_px = ImageReader(io.BytesIO(chart["png"])).getSize()
                width = min(page_width, 6.2 * inch)
                image = Image(io.BytesIO(chart["png"]), width=width, height=width * height_px / width_px)
                story.append(KeepTogether([image, Paragraph(esc(chart["title"]), small), Spacer(1, 8)]))
            except Exception:
                story.append(Paragraph(esc(f"The chart '{chart['title']}' could not be placed in the PDF."), small))
    else:
        story.append(Paragraph("No charts are included in this report.", small))

    # ---- recommendations --------------------------------------------------
    story.append(Paragraph("Recommendations", h2))
    recs = content["recommendations"]
    if recs["available"]:
        for _, rec in recs["table"].iterrows():
            block = [
                Paragraph(f"<b>{esc(rec['Priority'])} Priority</b> - Category: {esc(rec['Category'])}", body),
                Paragraph(f"<b>Issue:</b> {esc(rec['Issue'])}", body),
                Paragraph(f"<b>Finding:</b> {esc(rec['Finding'])}", body),
                Paragraph(f"<b>Recommendation:</b> {esc(rec['Recommendation'])}", body),
                Paragraph(f"<b>Suggested action:</b> {esc(rec['Suggested Action'])}", body),
                Spacer(1, 7),
            ]
            story.append(KeepTogether(block))
    else:
        story.append(Paragraph("Recommendation results are not currently available.", small))

    story.append(Paragraph("Recommended Actions", h2))
    for number, action in enumerate(content["actions"], start=1):
        story.append(Paragraph(esc(f"{number}. {action}"), body))
    if content["notes"]:
        story.append(Paragraph("Report Notes", h3))
        for note in content["notes"]:
            story.append(Paragraph(esc(note), small, bulletText="\u2022"))

    # "Page X of Y" needs the total page count, which is only known at the end,
    # so the canvas remembers every page and writes the footers when saving.
    class NumberedCanvas(canvas.Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._saved_pages = []

        def showPage(self):
            self._saved_pages.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._saved_pages)
            for state in self._saved_pages:
                self.__dict__.update(state)
                self.setFont("Helvetica", 8)
                self.setFillGray(0.4)
                self.drawCentredString(letter[0] / 2, 0.45 * inch, f"Page {self._pageNumber} of {total}")
                super().showPage()
            super().save()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=0.75 * inch, rightMargin=0.75 * inch,
                            topMargin=0.7 * inch, bottomMargin=0.8 * inch,
                            title=str(content["title"])[:200], author="Clarifile")
    doc.build(story, canvasmaker=NumberedCanvas)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# EXPORT DISPATCHER
# ---------------------------------------------------------------------------
def export_report(report_format, content, csv_table=None):
    """Create the chosen file. Returns (file_bytes, file_name, mime_type)."""
    choice = str(report_format).strip().lower()
    stem = _file_stem(content)
    if choice == "pdf":
        return generate_pdf_report(content), f"{stem}.pdf", "application/pdf"
    if choice in ("excel", "xlsx"):
        return (generate_excel_report(content), f"{stem}.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    if choice == "csv":
        table_name = csv_table or CSV_TABLE_OPTIONS[0]
        data = generate_csv_report(content, table_name)
        slug = re.sub(r"[^A-Za-z0-9]+", "_", table_name).strip("_").lower()
        return data, f"{stem}_{slug}.csv", "text/csv"
    raise ValueError(f"Unsupported report format: {report_format!r}. Choose one of {', '.join(REPORT_FORMATS)}.")
