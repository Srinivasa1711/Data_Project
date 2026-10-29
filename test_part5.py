"""
test_part5.py - runs without Streamlit. Uses clearly synthetic TEST FIXTURES
(these are test inputs, not report content) to exercise report_builder.py.
Run with: python test_part5.py
"""
import io

import numpy as np
import pandas as pd
from matplotlib.figure import Figure
from openpyxl import load_workbook
from pypdf import PdfReader

from report_builder import (CSV_TABLE_OPTIONS, export_report, generate_csv_report,
                            generate_excel_report, generate_pdf_report, generate_report)

rng = np.random.default_rng(0)
n = 200
df = pd.DataFrame({
    "age": rng.integers(18, 70, n).astype(float),
    "income": rng.normal(50000, 12000, n),
    "score": rng.normal(0, 1, n),
    "city": rng.choice(["Paris", "paris ", "Berlin", "=CMD()"], n),
    "num_as_text": rng.integers(1, 100, n).astype(str),
    "when": pd.date_range("2024-01-01", periods=n, freq="D"),
    "a<b & c": rng.normal(size=n),
})
df["income2"] = df["income"] * 2 + rng.normal(0, 100, n)  # strong correlation
df.loc[5:9, "age"] = np.nan
df.loc[0, "income"] = 900000  # outlier
df["income2"] = df["income"] * 2 + rng.normal(0, 100, n)  # recompute so the correlation stays strong
df = pd.concat([df, df.iloc[:3]], ignore_index=True)  # duplicates
snapshot = df.copy(deep=True)

fixture_quality = {"Quality Score": 78.5, "findings": ["Test finding from Part 2"]}
fig = Figure(figsize=(4, 3))
fig.subplots().plot([1, 2, 3])
fixture_analytics = {"findings": ["Test analytics insight"], "figures": {"Supplied test chart": fig}}
fixture_recs = [
    {"Priority": "low", "Category": "Types", "Issue": "Text numbers", "Finding": "f",
     "Recommendation": "r", "Suggested Action": "a"},
    {"priority": "High", "category": "Missing Data", "issue": "Missing ages", "finding": "5 missing",
     "recommendation": "Review", "suggested_action": "Investigate"},
]

content = generate_report(df, fixture_quality, fixture_analytics, fixture_recs, "My <Custom> Title & Co")
assert content["quality"]["score"] == "78.5"
assert content["recommendations"]["table"].loc[0, "Priority"] == "High"  # sorted
assert content["quality"]["duplicate_rows"] >= 3
assert len(content["quality"]["outlier_table"]) >= 1
assert len(content["analytics"]["strong_pairs"]) >= 1
assert any("Supplied test chart" in c["title"] for c in content["charts"])
assert "Top Value Share %" in content["analytics"]["categorical_summary"].columns
assert "Most Frequent Value" not in content["analytics"]["categorical_summary"].columns
print("content OK; charts:", [c["title"] for c in content["charts"]])

pdf = generate_pdf_report(report_content=content)
reader = PdfReader(io.BytesIO(pdf))
text = "".join(p.extract_text() for p in reader.pages)
assert pdf[:4] == b"%PDF" and len(reader.pages) >= 2
for needle in ("My <Custom> Title & Co", "Missing ages", "78.5", "Page 1 of", "Recommended Actions"):
    assert needle in text, needle
print("PDF OK:", len(reader.pages), "pages,", len(pdf), "bytes")

xlsx = generate_excel_report(report_content=content)
wb = load_workbook(io.BytesIO(xlsx))
print("Excel sheets:", wb.sheetnames)
assert wb.sheetnames[0] == "Summary" and "Recommendations" in wb.sheetnames
cells = [c.value for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
assert not any(v.startswith("=") for v in cells), "formula injection!"
assert wb["Recommendations"]["A2"].value == "High"

for name in CSV_TABLE_OPTIONS:
    try:
        data, fname, mime = export_report("CSV", content, name)
        print("CSV", name, "->", fname, len(data), "bytes")
    except ValueError as exc:
        print("CSV", name, "-> handled:", exc)
data, _, _ = export_report("CSV", content, "Data Types")
assert b"=CMD" not in data or b"'=CMD" in data

for bad in ("DOCX", ""):
    try:
        export_report(bad, content)
        raise SystemExit("invalid format not rejected")
    except ValueError:
        pass

for bad_df in (pd.DataFrame(), None, pd.DataFrame({"a": []})):
    try:
        generate_report(bad_df)
        raise SystemExit("empty df not rejected")
    except ValueError:
        pass

# Minimal inputs: no results at all, categorical only, numeric only, unhashable cells, duplicate names
for label, frame in {
    "categorical only": pd.DataFrame({"x": ["a", "b", None]}),
    "numeric only": pd.DataFrame({"x": [1.0, 2.0, 3.0, 4.0, 100.0]}),
    "unhashable": pd.DataFrame({"x": [[1], [2], [3]], "y": [1, 2, 3]}),
    "duplicate names": pd.DataFrame([[1, 2], [3, 4]], columns=["a", "a"]),
    "single row": pd.DataFrame({"x": [1], "y": ["z"]}),
}.items():
    c = generate_report(frame)
    generate_pdf_report(report_content=c)
    generate_excel_report(report_content=c)
    assert any("No recommendation" in note for note in c["notes"])
    assert c["quality"]["score"] is None
    print("minimal OK:", label)

assert df.equals(snapshot) and list(df.columns) == list(snapshot.columns), "DataFrame was modified!"
print("ALL TESTS PASSED (report_builder only; Streamlit UI not executed)")
